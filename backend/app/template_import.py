"""CSV / XLSX 模板批量导入。

流程：下载默认模板 → 上传 → 字段映射 → 校验预览（新增 / 更新 / 无变化 / 冲突 / 无效）
→ 确认提交 → 逐行结果 → 可整批撤销。

- 去重键是资产编号（必填），不按 IP 合并。同一文件里资产编号重复的行都标为冲突。
- 预览不访问设备、不改资产，只把逐行计划存成一个 preview 批次；提交时按预览时记下的
  资产修订号再核对一次，期间被别人改过的行标为冲突、不覆盖。
- 更新只改文件里填了值的列，空单元格保持原值。
- 撤销：只撤销导入后没有再被修改过的行（修订号仍等于导入后的值），被他人改过的保留并在结果里说明。
- 不执行公式：XLSX 里的公式单元格整行判为无效；导出结果时对 = + - @ 开头的值加前导单引号。
- 模板里不填密码，凭据用“凭据别名”引用已有凭据。
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from typing import Any

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.models import (
    Asset,
    AssetRelation,
    Credential,
    ImportBatch,
    ImportItem,
    ManagementEndpoint,
    Rack,
    Site,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
PREVIEW_TTL = timedelta(hours=24)


class ImportFileError(ValueError):
    pass


# ---- 模板定义 ----


@dataclass(frozen=True)
class Column:
    field: str
    header: str
    aliases: tuple[str, ...] = ()
    note: str = ""


COLUMNS: dict[str, Column] = {
    c.field: c
    for c in (
        Column("asset_no", "资产编号", ("编号", "asset_no", "asset number"), "必填，去重键；同一编号再次导入即更新"),
        Column("name", "名称", ("资产名称", "name"), "新增时必填"),
        Column("role", "类型", ("角色", "role", "type"), "服务器 / 交换机 / 路由器 / AP / 防火墙 / NAS / 存储 / UPS / 其他"),
        Column("form", "形态", ("form",), "物理机 / 虚拟机 / 容器"),
        Column("platform", "平台", ("系统", "platform", "os"), "如 debian、ubuntu、proxmox、openwrt、dell-idrac"),
        Column("vendor", "品牌", ("厂商", "vendor", "brand")),
        Column("model", "型号", ("model",)),
        Column("serial_no", "序列号", ("serial", "sn")),
        Column("address", "管理地址", ("ip", "地址", "address", "host")),
        Column("port", "端口", ("port",)),
        Column("protocol", "协议", ("protocol",), "ipmi / redfish / snmp / ssh / https / http / pve-api / esxi-api / winrm / other"),
        Column("site", "场所", ("位置", "site", "location"), "场所名称；不存在时自动新建"),
        Column("tags", "标签", ("tags",), "多个标签用逗号或分号分隔"),
        Column("collection_profile", "采集模板", ("collection_profile",), "P2 起生效"),
        Column("credential_alias", "凭据别名", ("凭据", "credential"), "引用已保存的凭据，不在模板里填密码"),
        Column("related_asset_no", "关联资产编号", ("宿主资产编号", "related", "host_asset_no"), "承载本资产的宿主（如虚拟机所在的物理机）；可引用同一文件里的资产"),
        Column("notes", "备注", ("notes", "说明")),
    )
}


@dataclass(frozen=True)
class TemplateKind:
    label: str
    columns: tuple[str, ...]
    form: str | None
    role: str | None
    protocol: str | None
    port: int | None
    platform: str = ""


TEMPLATES: dict[str, TemplateKind] = {
    "physical_server": TemplateKind(
        "物理服务器",
        ("asset_no", "name", "vendor", "model", "serial_no", "address", "port", "site", "tags", "collection_profile", "credential_alias", "related_asset_no", "notes"),
        "physical", "server", "ipmi", 623,
    ),
    "linux_server": TemplateKind(
        "Linux 服务器",
        ("asset_no", "name", "form", "platform", "address", "port", "site", "tags", "collection_profile", "credential_alias", "related_asset_no", "notes"),
        "vm", "server", "ssh", 22, "linux",
    ),
    "switch": TemplateKind(
        "交换机",
        ("asset_no", "name", "vendor", "model", "platform", "address", "port", "site", "tags", "collection_profile", "credential_alias", "related_asset_no", "notes"),
        "physical", "switch", "snmp", 161,
    ),
    "router": TemplateKind(
        "路由器",
        ("asset_no", "name", "vendor", "model", "platform", "address", "port", "site", "tags", "collection_profile", "credential_alias", "related_asset_no", "notes"),
        "physical", "router", "snmp", 161,
    ),
    "generic": TemplateKind(
        "通用资产",
        tuple(COLUMNS),
        None, None, None, None,
    ),
}

ROLE_VALUES = {
    "服务器": "server", "交换机": "switch", "路由器": "router", "ap": "ap", "无线ap": "ap", "防火墙": "firewall",
    "nas": "nas", "存储": "storage", "ups": "ups", "其他": "other",
    **{v: v for v in ("server", "switch", "router", "firewall", "storage", "other")},
}
FORM_VALUES = {"物理机": "physical", "物理": "physical", "物理设备": "physical", "虚拟机": "vm", "容器": "container",
               "physical": "physical", "vm": "vm", "container": "container"}
PROTOCOLS = {"ipmi", "redfish", "snmp", "ssh", "https", "http", "pve-api", "esxi-api", "winrm", "other"}


def _norm_header(text: str) -> str:
    return re.sub(r"[\s*＊（）()]", "", str(text)).lower()


def auto_mapping(headers: list[str]) -> dict[str, str | None]:
    lookup: dict[str, str] = {}
    for col in COLUMNS.values():
        for key in (col.header, col.field, *col.aliases):
            lookup[_norm_header(key)] = col.field
    return {h: lookup.get(_norm_header(h)) for h in headers}


# ---- 模板下载 ----


def template_headers(kind: str) -> list[str]:
    tpl = TEMPLATES[kind]
    required = {"asset_no", "name"} | ({"role"} if tpl.role is None else set())
    return [COLUMNS[f].header + ("*" if f in required else "") for f in tpl.columns]


def template_csv(kind: str) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerow(template_headers(kind))
    return ("﻿" + buf.getvalue()).encode("utf-8")


def template_xlsx(kind: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    tpl = TEMPLATES[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = tpl.label
    ws.append(template_headers(kind))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = 16
    ws.freeze_panes = "A2"
    guide = wb.create_sheet("填写说明")
    guide.append(["列", "说明"])
    for f in tpl.columns:
        guide.append([COLUMNS[f].header, COLUMNS[f].note])
    defaults = []
    if tpl.form:
        defaults.append(f"形态默认 {tpl.form}")
    if tpl.role:
        defaults.append(f"类型默认 {tpl.role}")
    if tpl.protocol:
        defaults.append(f"管理地址按 {tpl.protocol} 协议、默认端口 {tpl.port}")
    guide.append([])
    guide.append(["本模板", "；".join(defaults) or "类型必填；填写管理地址时协议必填"])
    guide.append(["注意", "第一个工作表的第一行是表头；不要在单元格里写公式，含公式的行会被判为无效；不在模板里填写密码。"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---- 读取上传文件 ----


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[tuple[int, dict[str, str]]]
    formula_rows: set[int] = field(default_factory=set)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def parse_file(filename: str, content: bytes) -> ParsedFile:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(f"文件超过 {MAX_FILE_BYTES // 1024 // 1024} MB 上限")
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        parsed = _parse_xlsx(content)
    elif lower.endswith(".csv"):
        parsed = _parse_csv(content)
    else:
        raise ImportFileError("只支持 .csv 与 .xlsx 文件")
    if not parsed.headers:
        raise ImportFileError("文件里没有表头")
    if len(parsed.rows) > MAX_ROWS:
        raise ImportFileError(f"数据行超过 {MAX_ROWS} 行上限，请拆分后分批导入")
    return parsed


def _parse_csv(content: bytes) -> ParsedFile:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError("无法识别文件编码，请另存为 UTF-8 CSV")
    reader = csv.reader(io.StringIO(text))
    rows = [r for r in reader]
    while rows and not any(c.strip() for c in rows[0]):
        rows.pop(0)
    if not rows:
        return ParsedFile([], [])
    headers = [h.strip() for h in rows[0]]
    data = []
    for i, r in enumerate(rows[1:], start=2):
        if any(c.strip() for c in r):
            data.append((i, {headers[j]: r[j].strip() for j in range(min(len(headers), len(r))) if headers[j]}))
    return ParsedFile(headers, data)


def _parse_xlsx(content: bytes) -> ParsedFile:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError("无法打开 XLSX 文件") from exc
    try:
        ws = wb.worksheets[0]
        headers: list[str] = []
        data: list[tuple[int, dict[str, str]]] = []
        formulas: set[int] = set()
        for row in ws.iter_rows():
            row_no = row[0].row if row and hasattr(row[0], "row") else 0
            values = [_cell_text(getattr(c, "value", None)) for c in row]
            if not headers:
                if any(values):
                    headers = values
                continue
            if not any(values):
                continue
            if any(getattr(c, "data_type", "") == "f" or str(getattr(c, "value", "") or "").startswith("=") for c in row):
                formulas.add(row_no)
            data.append((row_no, {headers[j]: values[j] for j in range(min(len(headers), len(values))) if headers[j]}))
            if len(data) > MAX_ROWS:
                break
        return ParsedFile([h for h in headers], data, formulas)
    finally:
        wb.close()


# ---- 预览 ----


ASSET_FIELDS = ("name", "role", "form", "platform", "vendor", "model", "serial_no", "collection_profile", "notes", "tags", "site_id")


@dataclass
class RowPlan:
    row_no: int
    asset_no: str
    name: str
    action: str
    messages: list[str] = field(default_factory=list)
    data: dict[str, Any] = field(default_factory=dict)
    asset_id: Any = None
    base_revision: int | None = None


def _split_tags(text: str) -> list[str]:
    return [t.strip() for t in re.split(r"[,，;；、]", text) if t.strip()][:50]


def build_preview(
    db: Session,
    parsed: ParsedFile,
    kind: str,
    mapping: dict[str, str | None] | None = None,
) -> tuple[dict[str, str | None], list[RowPlan]]:
    tpl = TEMPLATES[kind]
    mapping = {**auto_mapping(parsed.headers), **(mapping or {})}
    used = [f for f in mapping.values() if f]
    dupes = {f for f in used if used.count(f) > 1}
    if dupes:
        raise ImportFileError("多列映射到了同一字段：" + "、".join(COLUMNS[f].header for f in dupes if f in COLUMNS))
    if "asset_no" not in used:
        raise ImportFileError("缺少“资产编号”列：它是去重键，必须提供")

    site_ids: dict[str, list[Any]] = {}
    for s in db.scalars(select(Site)):
        site_ids.setdefault(s.name, []).append(s.id)
    cred_ids = {c.alias: c.id for c in db.scalars(select(Credential))}

    raw_rows: list[tuple[int, dict[str, str]]] = []
    for row_no, cells in parsed.rows:
        values: dict[str, str] = {}
        for header, value in cells.items():
            f = mapping.get(header)
            if f:
                values[f] = value
        raw_rows.append((row_no, values))

    counts: dict[str, int] = {}
    for _, values in raw_rows:
        no = values.get("asset_no", "")
        if no:
            counts[no] = counts.get(no, 0) + 1
    file_nos = set(counts)

    plans: list[RowPlan] = []
    for row_no, v in raw_rows:
        asset_no = v.get("asset_no", "")[:100]
        plan = RowPlan(row_no=row_no, asset_no=asset_no, name=v.get("name", "")[:200], action="invalid")
        errors: list[str] = []
        notes: list[str] = []
        if row_no in parsed.formula_rows:
            errors.append("包含公式，不执行公式，请改成纯文本或数值")
        if not asset_no:
            errors.append("资产编号必填")
        elif counts.get(asset_no, 0) > 1:
            plan.action = "conflict"
            plan.messages = ["同一文件中资产编号重复"]
            plans.append(plan)
            continue

        existing = db.scalar(select(Asset).where(Asset.asset_no == asset_no)) if asset_no else None
        if existing is not None and existing.deleted_at is not None:
            # 资产编号唯一：回收站里的资产仍占着编号，导入不悄悄覆盖或恢复它
            plan.action = "conflict"
            plan.messages = [f"资产编号与回收站里的资产“{existing.name}”相同，请先在回收站恢复或彻底删除"]
            plans.append(plan)
            continue
        data: dict[str, Any] = {}

        if v.get("role"):
            role = ROLE_VALUES.get(v["role"].strip().lower())
            if role is None:
                errors.append(f"类型“{v['role']}”无法识别")
            else:
                data["role"] = role
        elif existing is None:
            if tpl.role:
                data["role"] = tpl.role
            else:
                errors.append("类型必填")
        if v.get("form"):
            form = FORM_VALUES.get(v["form"].strip().lower())
            if form is None:
                errors.append(f"形态“{v['form']}”无法识别")
            else:
                data["form"] = form
        elif existing is None:
            data["form"] = tpl.form or "physical"
        if existing is None and not v.get("name"):
            errors.append("名称必填")
        for f, limit in (("name", 200), ("platform", 64), ("vendor", 100), ("model", 100), ("serial_no", 100), ("collection_profile", 100)):
            if v.get(f):
                if len(v[f]) > limit:
                    errors.append(f"{COLUMNS[f].header}超过 {limit} 个字符")
                data[f] = v[f][:limit]
        if existing is None and "platform" not in data and tpl.platform:
            data["platform"] = tpl.platform
        if v.get("notes"):
            data["notes"] = v["notes"]
        if v.get("tags"):
            data["tags"] = _split_tags(v["tags"])
        if v.get("collection_profile"):
            notes.append("采集模板在 P2 生效，先保存名称")

        if v.get("site"):
            ids = site_ids.get(v["site"], [])
            if len(ids) > 1:
                errors.append(f"场所“{v['site']}”有多个同名，无法确定")
            elif ids:
                data["site_id"] = str(ids[0])
            else:
                data["new_site"] = v["site"][:200]
                notes.append(f"将新建场所“{v['site']}”")

        address = v.get("address", "")
        if address:
            protocol = (v.get("protocol") or tpl.protocol or "").strip().lower()
            if not protocol:
                errors.append("填写管理地址时协议必填")
            elif protocol not in PROTOCOLS:
                errors.append(f"协议“{protocol}”无法识别")
            port: int | None = tpl.port if not v.get("protocol") else None
            if v.get("port"):
                try:
                    port = int(v["port"])
                    if not 1 <= port <= 65535:
                        raise ValueError
                except ValueError:
                    errors.append(f"端口“{v['port']}”不是 1–65535 的整数")
            data["endpoint"] = {"address": address[:255], "protocol": protocol, "port": port}
            if v.get("credential_alias"):
                cid = cred_ids.get(v["credential_alias"])
                if cid is None:
                    errors.append(f"凭据别名“{v['credential_alias']}”不存在，请先在凭据管理里创建")
                else:
                    data["endpoint"]["credential_id"] = str(cid)
        elif v.get("credential_alias") or v.get("port"):
            errors.append("填写了端口或凭据别名但没有管理地址")

        rel = v.get("related_asset_no", "")
        if rel:
            if rel == asset_no:
                errors.append("关联资产编号不能是自己")
            elif rel not in file_nos and db.scalar(select(Asset.id).where(Asset.asset_no == rel, Asset.deleted_at.is_(None))) is None:
                errors.append(f"关联资产编号“{rel}”在库里和本文件里都不存在")
            else:
                data["related_asset_no"] = rel

        plan.data = data
        if errors:
            plan.messages = errors
            plans.append(plan)
            continue
        plan.messages = notes
        if existing is None:
            plan.action = "create"
        else:
            plan.asset_id = existing.id
            plan.base_revision = existing.revision
            plan.name = plan.name or existing.name
            changes = _diff(db, existing, data)
            plan.action = "update" if changes else "unchanged"
            if changes:
                plan.messages = [*notes, "更新：" + "、".join(changes)]
        plans.append(plan)
    return mapping, plans


def _diff(db: Session, asset: Asset, data: dict[str, Any]) -> list[str]:
    labels = {"site_id": "场所", **{c.field: c.header for c in COLUMNS.values()}}
    changed = []
    for f in ASSET_FIELDS:
        if f in data:
            current = str(getattr(asset, f)) if f == "site_id" and asset.site_id else getattr(asset, f)
            if current != data[f]:
                changed.append(labels[f])
    if "new_site" in data:
        changed.append("场所")
    ep = data.get("endpoint")
    if ep:
        cur = _match_endpoint(db, asset, ep["protocol"])
        if cur is None or cur.address != ep["address"] or (ep.get("port") and cur.port != ep["port"]) or (
            ep.get("credential_id") and str(cur.credential_id) != ep["credential_id"]
        ):
            changed.append("管理端点")
    rel = data.get("related_asset_no")
    if rel:
        host = db.scalar(select(Asset).where(Asset.asset_no == rel, Asset.deleted_at.is_(None)))
        if host is None or db.scalar(
            select(AssetRelation).where(
                AssetRelation.source_asset_id == host.id, AssetRelation.target_asset_id == asset.id, AssetRelation.kind == "hosts"
            )
        ) is None:
            changed.append("关联资产")
    return changed


def _match_endpoint(db: Session, asset: Asset, protocol: str) -> ManagementEndpoint | None:
    return db.scalar(
        select(ManagementEndpoint)
        .where(ManagementEndpoint.asset_id == asset.id, ManagementEndpoint.protocol == protocol)
        .order_by(ManagementEndpoint.created_at)
    )


def save_preview(
    db: Session, kind: str, filename: str, actor: str, mapping: dict[str, str | None], plans: list[RowPlan]
) -> ImportBatch:
    db.execute(
        delete(ImportBatch).where(ImportBatch.status == "preview", ImportBatch.created_at < datetime.now(UTC) - PREVIEW_TTL)
    )
    batch = ImportBatch(source="template", template_kind=kind, filename=filename[:255], created_by=actor, status="preview")
    for p in plans:
        batch.items.append(
            ImportItem(
                row_no=p.row_no, asset_no=p.asset_no, name=p.name, action=p.action, messages=p.messages,
                data=p.data, asset_id=p.asset_id, base_revision=p.base_revision,
            )
        )
    batch.summary = {"mapping": mapping, "preview": summarize(batch.items)}
    db.add(batch)
    db.flush()
    return batch


def summarize(items: list[ImportItem], attr: str = "action") -> dict[str, int]:
    out: dict[str, int] = {}
    for item in items:
        key = getattr(item, attr)
        if key:
            out[key] = out.get(key, 0) + 1
    return out


# ---- 提交 ----


def _snapshot(asset: Asset) -> dict[str, Any]:
    return {f: (str(asset.site_id) if f == "site_id" and asset.site_id else getattr(asset, f)) for f in ASSET_FIELDS}


def _apply_fields(asset: Asset, data: dict[str, Any], site_lookup: dict[str, Any]) -> None:
    import uuid as _uuid

    for f in ASSET_FIELDS:
        if f in data:
            value = data[f]
            if f == "site_id" and value is not None:
                value = _uuid.UUID(value)
            setattr(asset, f, value)
    if "new_site" in data:
        asset.site_id = site_lookup[data["new_site"]]


def commit_batch(db: Session, batch: ImportBatch) -> dict[str, int]:
    import uuid as _uuid

    if batch.status != "preview":
        raise ImportFileError("该批次已提交或已撤销")
    created_sites: dict[str, Any] = {}

    def site_for(name: str) -> Any:
        if name not in created_sites:
            existing = db.scalar(select(Site).where(Site.name == name))
            if existing is not None:
                created_sites[name] = existing.id
            else:
                site = Site(name=name, kind="site")
                db.add(site)
                db.flush()
                created_sites[name] = site.id
                batch.summary = {**batch.summary, "created_sites": [*batch.summary.get("created_sites", []), str(site.id)]}
        return created_sites[name]

    by_no: dict[str, Asset] = {}
    for item in batch.items:
        data = item.data
        if item.action not in ("create", "update"):
            item.action = {"unchanged": "unchanged", "invalid": "skipped", "conflict": "conflict"}[item.action]
            continue
        if "new_site" in data:
            site_for(data["new_site"])
        if item.action == "create":
            if db.scalar(select(Asset.id).where(Asset.asset_no == item.asset_no)) is not None:
                item.action = "conflict"
                item.messages = [*item.messages, "预览之后已有人创建了同编号资产，未导入"]
                continue
            asset = Asset(asset_no=item.asset_no, name=data["name"], form=data["form"], role=data["role"])
            _apply_fields(asset, data, created_sites)
            db.add(asset)
            db.flush()
            refs: dict[str, Any] = {}
            ep = data.get("endpoint")
            if ep:
                e = ManagementEndpoint(
                    asset_id=asset.id, address=ep["address"], protocol=ep["protocol"], port=ep.get("port"),
                    credential_id=_uuid.UUID(ep["credential_id"]) if ep.get("credential_id") else None,
                )
                db.add(e)
            item.asset_id, item.after_revision, item.action, item.created_refs = asset.id, asset.revision, "created", refs
        else:
            asset = db.get(Asset, item.asset_id) if item.asset_id else None
            if asset is None or asset.revision != item.base_revision:
                item.action = "conflict"
                item.messages = [*item.messages, "预览之后该资产已被修改或删除，未覆盖"]
                continue
            before = _snapshot(asset)
            refs = {}
            _apply_fields(asset, data, created_sites)
            ep = data.get("endpoint")
            if ep:
                cur = _match_endpoint(db, asset, ep["protocol"])
                if cur is None:
                    e = ManagementEndpoint(
                        asset_id=asset.id, address=ep["address"], protocol=ep["protocol"], port=ep.get("port"),
                        credential_id=_uuid.UUID(ep["credential_id"]) if ep.get("credential_id") else None,
                    )
                    db.add(e)
                    db.flush()
                    refs["endpoint_created"] = str(e.id)
                else:
                    before["endpoint"] = {"id": str(cur.id), "address": cur.address, "port": cur.port,
                                          "credential_id": str(cur.credential_id) if cur.credential_id else None}
                    cur.address = ep["address"]
                    if ep.get("port"):
                        cur.port = ep["port"]
                    if ep.get("credential_id"):
                        cur.credential_id = _uuid.UUID(ep["credential_id"])
            asset.revision += 1
            item.before, item.created_refs, item.after_revision, item.action = before, refs, asset.revision, "updated"
        by_no[item.asset_no] = asset

    db.flush()
    # 第二遍建立“承载”关系，这样同一文件里先后出现的宿主和虚拟机都能关联上
    for item in batch.items:
        rel_no = item.data.get("related_asset_no")
        if item.action not in ("created", "updated") or not rel_no:
            continue
        host = by_no.get(rel_no) or db.scalar(select(Asset).where(Asset.asset_no == rel_no, Asset.deleted_at.is_(None)))
        asset = by_no[item.asset_no]
        if host is None:
            item.messages = [*item.messages, f"宿主 {rel_no} 未导入，未建立关联"]
            continue
        exists = db.scalar(
            select(AssetRelation).where(
                AssetRelation.source_asset_id == host.id, AssetRelation.target_asset_id == asset.id, AssetRelation.kind == "hosts"
            )
        )
        if exists is None:
            rel = AssetRelation(source_asset_id=host.id, target_asset_id=asset.id, kind="hosts", origin="import")
            db.add(rel)
            db.flush()
            item.created_refs = {**item.created_refs, "relation": str(rel.id)}

    batch.status = "applied"
    batch.applied_at = datetime.now(UTC)
    result = summarize(batch.items)
    batch.summary = {**batch.summary, "result": result}
    return result


# ---- 撤销 ----


def undo_batch(db: Session, batch: ImportBatch) -> dict[str, int]:
    import uuid as _uuid

    if batch.status != "applied":
        raise ImportFileError("只有已提交的批次可以撤销")
    for item in reversed(batch.items):
        if item.action not in ("created", "updated"):
            continue
        rel_id = item.created_refs.get("relation")
        if rel_id:
            db.execute(delete(AssetRelation).where(AssetRelation.id == _uuid.UUID(rel_id)))
        asset = db.get(Asset, item.asset_id) if item.asset_id else None
        if asset is None:
            item.undo_result = "gone"
            continue
        if asset.revision != item.after_revision:
            item.undo_result = "kept_modified"
            continue
        if item.action == "created":
            db.delete(asset)
            item.undo_result = "removed"
            continue
        before = dict(item.before)
        ep_before = before.pop("endpoint", None)
        _apply_fields(asset, before, {})
        if ep_before:
            ep = db.get(ManagementEndpoint, _uuid.UUID(ep_before["id"]))
            if ep is not None:
                ep.address, ep.port = ep_before["address"], ep_before["port"]
                ep.credential_id = _uuid.UUID(ep_before["credential_id"]) if ep_before["credential_id"] else None
        created_ep = item.created_refs.get("endpoint_created")
        if created_ep:
            db.execute(delete(ManagementEndpoint).where(ManagementEndpoint.id == _uuid.UUID(created_ep)))
        asset.revision += 1
        item.undo_result = "restored"
    db.flush()
    for sid in batch.summary.get("created_sites", []):
        site_id = _uuid.UUID(sid)
        used = (
            db.scalar(select(func.count()).select_from(Asset).where(Asset.site_id == site_id))
            or db.scalar(select(func.count()).select_from(Rack).where(Rack.site_id == site_id))
            or db.scalar(select(func.count()).select_from(Site).where(Site.parent_id == site_id))
        )
        if not used:
            db.execute(delete(Site).where(Site.id == site_id))
    batch.status = "undone"
    batch.undone_at = datetime.now(UTC)
    result = summarize(batch.items, "undo_result")
    batch.summary = {**batch.summary, "undo": result}
    return result


# ---- 结果导出 ----


def _safe_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


ACTION_LABELS = {
    "create": "新增", "update": "更新", "unchanged": "无变化", "invalid": "无效", "conflict": "冲突",
    "created": "已新增", "updated": "已更新", "skipped": "已跳过",
    "removed": "已撤销（删除）", "restored": "已撤销（恢复原值）", "kept_modified": "未撤销（导入后已被修改）", "gone": "资产已不存在",
}


def results_csv(batch: ImportBatch) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["行号", "资产编号", "名称", "结果", "撤销结果", "说明"])
    for item in batch.items:
        w.writerow([
            item.row_no, _safe_cell(item.asset_no), _safe_cell(item.name), ACTION_LABELS.get(item.action, item.action),
            ACTION_LABELS.get(item.undo_result, item.undo_result), _safe_cell("；".join(item.messages)),
        ])
    return ("﻿" + buf.getvalue()).encode("utf-8")
