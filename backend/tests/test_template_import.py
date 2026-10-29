from __future__ import annotations

import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.template_import import TEMPLATES, template_headers
from tests.api_helpers import add_user, login


@pytest.fixture
def op(app_factory, factory) -> TestClient:  # type: ignore[no-untyped-def]
    add_user(factory, "olivia", "operator")
    return login(app_factory(), "olivia")


def csv_bytes(rows: list[list[str]], encoding: str = "utf-8-sig") -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode(encoding)


def preview(client: TestClient, content: bytes, kind: str = "physical_server", filename: str = "a.csv", mapping: str | None = None):  # type: ignore[no-untyped-def]
    data = {"kind": kind}
    if mapping:
        data["mapping"] = mapping
    return client.post("/api/import/preview", data=data, files={"file": (filename, content)})


def actions(batch: dict) -> dict[str, str]:  # type: ignore[type-arg]
    return {i["asset_no"]: i["action"] for i in batch["items"]}


HEAD = ["资产编号*", "名称*", "品牌", "型号", "管理地址", "端口", "场所", "标签", "关联资产编号"]


def test_templates_download(op: TestClient) -> None:
    kinds = [t["kind"] for t in op.get("/api/import/templates").json()]
    assert kinds == list(TEMPLATES)
    resp = op.get("/api/import/templates/physical_server.csv")
    assert resp.status_code == 200
    assert resp.content.decode("utf-8-sig").strip() == ",".join(template_headers("physical_server"))
    assert "密码" not in resp.content.decode("utf-8-sig")
    x = op.get("/api/import/templates/generic.xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert [c.value for c in wb.worksheets[0][1]] == template_headers("generic")
    assert wb.sheetnames[1] == "填写说明"
    assert op.get("/api/import/templates/nope.csv").status_code == 404


def test_preview_commit_and_reimport(op: TestClient) -> None:
    rows = [HEAD,
            ["S-1", "R740 计算", "Dell", "R740", "192.0.2.11", "", "机房A", "生产,计算", ""],
            ["S-2", "R750 存储", "Dell", "R750", "192.0.2.12", "6230", "机房A", "", "S-1"]]
    pv = preview(op, csv_bytes(rows))
    assert pv.status_code == 201, pv.text
    batch = pv.json()
    assert batch["status"] == "preview"
    assert actions(batch) == {"S-1": "create", "S-2": "create"}
    assert any("将新建场所" in m for m in batch["items"][0]["messages"])
    # 预览不改资产
    assert op.get("/api/assets").json()["total"] == 0

    done = op.post(f"/api/import/{batch['id']}/commit").json()
    assert actions(done) == {"S-1": "created", "S-2": "created"}
    assets = {a["asset_no"]: a for a in op.get("/api/assets").json()["items"]}
    assert assets["S-1"]["form"] == "physical" and assets["S-1"]["role"] == "server"
    assert assets["S-1"]["tags"] == ["生产", "计算"]
    assert assets["S-1"]["endpoints"][0]["port"] == 623
    assert assets["S-2"]["endpoints"][0]["port"] == 6230
    assert len(op.get("/api/sites").json()) == 1
    rel = op.get(f"/api/assets/{assets['S-2']['id']}").json()["relations"]
    assert rel[0]["source_asset_id"] == assets["S-1"]["id"] and rel[0]["kind"] == "hosts"

    # 同一文件再导入一次：全部无变化，不生成重复资产
    again = preview(op, csv_bytes(rows)).json()
    assert actions(again) == {"S-1": "unchanged", "S-2": "unchanged"}
    op.post(f"/api/import/{again['id']}/commit")
    assert op.get("/api/assets").json()["total"] == 2

    # 改一列：只更新该资产；空单元格不清空原值
    rows2 = [HEAD, ["S-1", "", "", "R740xd", "", "", "", "", ""]]
    pv2 = preview(op, csv_bytes(rows2)).json()
    assert actions(pv2) == {"S-1": "update"}
    op.post(f"/api/import/{pv2['id']}/commit")
    s1 = op.get(f"/api/assets/{assets['S-1']['id']}").json()
    assert s1["model"] == "R740xd" and s1["name"] == "R740 计算" and s1["vendor"] == "Dell"


def test_invalid_rows_reported(op: TestClient) -> None:
    rows = [HEAD + ["凭据别名"],
            ["", "无编号", "", "", "", "", "", "", "", ""],
            ["D-1", "重复甲", "", "", "", "", "", "", "", ""],
            ["D-1", "重复乙", "", "", "", "", "", "", "", ""],
            ["E-1", "端口错", "", "", "192.0.2.1", "99999", "", "", "", ""],
            ["E-2", "没有宿主", "", "", "", "", "", "", "X-404", ""],
            ["E-3", "凭据不存在", "", "", "192.0.2.3", "", "", "", "", "no-such-alias"],
            ["E-4", "", "", "", "", "", "", "", "", ""],
            ["OK-1", "正常", "", "", "", "", "", "", "", ""]]
    batch = preview(op, csv_bytes(rows)).json()
    by_row = {i["row_no"]: i for i in batch["items"]}
    assert by_row[2]["action"] == "invalid" and "资产编号必填" in by_row[2]["messages"]
    assert by_row[3]["action"] == by_row[4]["action"] == "conflict"
    assert by_row[5]["action"] == "invalid"
    assert by_row[6]["action"] == "invalid"
    assert by_row[7]["action"] == "invalid"
    assert by_row[8]["action"] == "invalid"  # 新增时名称必填
    assert by_row[9]["action"] == "create"
    done = op.post(f"/api/import/{batch['id']}/commit").json()
    assert sorted(i["action"] for i in done["items"]) == ["conflict", "conflict", "created", "skipped", "skipped", "skipped", "skipped", "skipped"]
    assert op.get("/api/assets").json()["total"] == 1


def test_undo_skips_assets_modified_after_import(op: TestClient) -> None:
    existing = op.post("/api/assets", json={"asset_no": "U-0", "name": "原有", "form": "physical", "role": "server", "model": "旧型号"}).json()
    rows = [HEAD, ["U-0", "", "", "新型号", "", "", "", "", ""], ["U-1", "新甲", "", "", "", "", "", "", ""], ["U-2", "新乙", "", "", "", "", "", "", ""]]
    batch = preview(op, csv_bytes(rows)).json()
    op.post(f"/api/import/{batch['id']}/commit")

    # 导入后有人改了 U-2：撤销时保留它
    u2 = next(a for a in op.get("/api/assets").json()["items"] if a["asset_no"] == "U-2")
    op.patch(f"/api/assets/{u2['id']}", json={"revision": u2["revision"], "notes": "人工补充"})

    undone = op.post(f"/api/import/{batch['id']}/undo").json()
    assert undone["status"] == "undone"
    res = {i["asset_no"]: i["undo_result"] for i in undone["items"]}
    assert res == {"U-0": "restored", "U-1": "removed", "U-2": "kept_modified"}
    nos = {a["asset_no"]: a for a in op.get("/api/assets").json()["items"]}
    assert set(nos) == {"U-0", "U-2"}
    assert nos["U-0"]["model"] == "旧型号"
    assert op.get(f"/api/assets/{existing['id']}").json()["model"] == "旧型号"
    assert op.post(f"/api/import/{batch['id']}/undo").status_code == 409


def test_commit_detects_change_after_preview(op: TestClient) -> None:
    a = op.post("/api/assets", json={"asset_no": "C-1", "name": "甲", "form": "physical", "role": "server"}).json()
    batch = preview(op, csv_bytes([HEAD, ["C-1", "", "", "R640", "", "", "", "", ""]])).json()
    op.patch(f"/api/assets/{a['id']}", json={"revision": 1, "name": "被人改了"})
    done = op.post(f"/api/import/{batch['id']}/commit").json()
    assert done["items"][0]["action"] == "conflict"
    assert op.get(f"/api/assets/{a['id']}").json()["model"] == ""
    assert op.post(f"/api/import/{batch['id']}/commit").status_code == 409


def test_xlsx_formula_rows_rejected(op: TestClient) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["资产编号", "名称", "类型"])
    ws.append(["X-1", "正常", "交换机"])
    ws.append(["X-2", "=1+1", "交换机"])
    out = io.BytesIO()
    wb.save(out)
    batch = preview(op, out.getvalue(), kind="generic", filename="t.xlsx").json()
    got = {i["asset_no"]: i for i in batch["items"]}
    assert got["X-1"]["action"] == "create"
    assert got["X-2"]["action"] == "invalid" and "公式" in got["X-2"]["messages"][0]


def test_generic_requires_role_and_protocol(op: TestClient) -> None:
    rows = [["资产编号", "名称", "类型", "管理地址", "协议"],
            ["G-1", "无类型", "", "", ""],
            ["G-2", "无协议", "交换机", "192.0.2.5", ""],
            ["G-3", "防火墙", "防火墙", "192.0.2.6", "https"]]
    batch = preview(op, csv_bytes(rows), kind="generic").json()
    assert actions(batch) == {"G-1": "invalid", "G-2": "invalid", "G-3": "create"}


def test_gbk_csv_and_custom_mapping(op: TestClient) -> None:
    rows = [["编号", "主机名", "IP"], ["M-1", "老表格", "192.0.2.9"]]
    batch = preview(op, csv_bytes(rows, "gbk"), kind="linux_server", mapping='{"主机名": "name"}').json()
    assert batch["summary"]["mapping"] == {"编号": "asset_no", "主机名": "name", "IP": "address"}
    assert actions(batch) == {"M-1": "create"}
    op.post(f"/api/import/{batch['id']}/commit")
    a = op.get("/api/assets").json()["items"][0]
    assert (a["form"], a["platform"], a["endpoints"][0]["protocol"], a["endpoints"][0]["port"]) == ("vm", "linux", "ssh", 22)


def test_missing_asset_no_column_and_limits(op: TestClient) -> None:
    assert preview(op, csv_bytes([["名称"], ["甲"]])).status_code == 422
    assert preview(op, b"x", filename="a.txt").status_code == 422
    big = csv_bytes([["资产编号", "名称"]] + [[f"N-{i}", "x"] for i in range(5001)])
    assert preview(op, big).status_code == 422


def test_results_csv_escapes_formulas(op: TestClient) -> None:
    batch = preview(op, csv_bytes([HEAD, ["=HYPERLINK(1)", "@cmd", "", "", "", "", "", "", ""]])).json()
    text = op.get(f"/api/import/{batch['id']}/results.csv").content.decode("utf-8-sig")
    assert "'=HYPERLINK(1)" in text and "'@cmd" in text


def test_viewer_cannot_import(app_factory, factory) -> None:  # type: ignore[no-untyped-def]
    add_user(factory, "victor", "viewer")
    v = login(app_factory(), "victor")
    assert v.get("/api/import/templates/router.csv").status_code == 200
    assert preview(v, csv_bytes([HEAD])).status_code == 403
